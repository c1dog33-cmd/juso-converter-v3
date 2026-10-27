import io
import re
import urllib.parse
import openpyxl
import pandas as pd
import requests
import streamlit as st
from openpyxl.styles import Font

# 🚀 행정안전부 도로명주소 오픈 API 승인키 적용 완료
JUSO_CONFIRM_KEY = "U01TX0FVVEgyMDI2MTAyMzE3MDUxMjEyMDgzMzI="  

def fix_zipcode(val):
    if pd.isna(val) or val is None:
        return ""
    val_str = str(val).split('.')[0].strip()
    return val_str.zfill(5) if val_str else ""

def remove_illegal_chars(val):
    if isinstance(val, str):
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', val)
    return val

def remove_duplicate_parentheses(addr_str):
    if not addr_str:
        return addr_str
    parts = re.split(r'(\([^)]+\))', addr_str)
    seen = set()
    new_parts = []
    for part in parts:
        if part.startswith('(') and part.endswith(')'):
            normalized = re.sub(r'\s+', '', part)
            if normalized in seen:
                continue
            seen.add(normalized)
            new_parts.append(part)
        else:
            new_parts.append(part)
    return ''.join(new_parts)

def format_dong_ho(addr_str):
    if not addr_str:
        return addr_str
    
    # '106-1001호' 형태를 '106동 1001호'로 자동 변환
    addr_str = re.sub(r'\b(\d+)-(\d+호)\b', r'\1동 \2', str(addr_str))
    
    # '107동109호' 처럼 동과 호수 사이에 공백이 없는 경우 자동 띄어쓰기 삽입
    addr_str = re.sub(r'(\d+동)(\d+호)', r'\1 \2', addr_str)
    addr_str = re.sub(r'([A-Z가-힣]\s*동)(\d+호)', r'\1 \2', addr_str)
    return addr_str

def clean_product_name(val):
    if pd.isna(val) or not str(val).strip():
        return val
    s = str(val).strip()
    
    has_film = False
    if ('풀액정필름2장' in s or '추가액정필름' in s) and '필요없음' not in s and ('풀액정' in s or '액정필름' in s):
        has_film = True

    target_texts = [s]
    if '//' in s: target_texts.insert(0, s.split('//')[-1])
    if '모델선택:' in s: target_texts.insert(0, s.split('모델선택:')[-1])
    if '/' in s: target_texts.insert(0, s.split('/')[-1])

    smartphone_match = None
    for t in target_texts:
        smartphone_match = re.search(r'(S\s*10\s*5G|S\s*10\s*플러스|S\s*20\s*플러스|S\s*20\s*울트라|S\s*22\s*플러스|S\s*22\s*울트라|노트\s*10\s*플러스|노트\s*20\s*울트라|노트\s*20\s*플러스|노트\s*\d+\s*울트라|노트\s*\d+\s*플러스|노트\s*\d+|S\s*\d+(?:\s*플러스|\s*울트라|\s*5G|\s*엣지)?|A\s*\d+|Z\s*(?:플립|폴드)\s*\d*|와이드\s*\d+|N\s*\d+|V\s*50S|V\s*50|V\s*\d+|벨벳|윙)', t, re.IGNORECASE)
        if smartphone_match: break

    if not smartphone_match:
        smartphone_match = re.search(r'(S\s*10\s*5G|S\s*10\s*플러스|S\s*20\s*플러스|S\s*20\s*울트라|S\s*22\s*플러스|S\s*22\s*울트라|노트\s*10\s*플러스|노트\s*20\s*울트라|노트\s*20\s*플러스|노트\s*\d+\s*울트라|노트\s*\d+\s*플러스|노트\s*\d+|S\s*\d+(?:\s*플러스|\s*울트라|\s*5G|\s*엣지)?|A\s*\d+|Z\s*(?:플립|폴드)\s*\d*|와이드\s*\d+|N\s*\d+|V\s*50S|V\s*50|V\s*\d+|벨벳|윙)', s, re.IGNORECASE)

    if smartphone_match:
        raw_model = smartphone_match.group(1).replace(" ", "").upper()
        if raw_model == 'S105G': target_model = 'S10 5G'
        elif raw_model == 'S10플러스': target_model = 'S10플러스'
        elif raw_model == 'S20플러스': target_model = 'S20플러스'
        elif raw_model == 'S22플러스': target_model = 'S22플러스'
        elif raw_model == '노트10플러스': target_model = '노트10플러스'
        elif raw_model in ['N960', 'SM-N960']: target_model = '노트9'
        elif raw_model == 'V50S': target_model = 'V50S'
        elif raw_model == 'V50': target_model = 'V50'
        elif raw_model == '노트20울트라': target_model = '노트20울트라'
        elif raw_model == '와이드5': target_model = 'A22'
        elif raw_model == '와이드6': target_model = 'A13'
        else: target_model = raw_model

        qty = 1
        if '☞' in s:
            parts = s.split('☞')
            after = parts[1].strip() if len(parts) > 1 else ""
            qty_match = re.match(r'^(\d+)\s*(EA|개)?', after, re.IGNORECASE)
            if qty_match: qty = int(qty_match.group(1))
        
        res = f"{target_model}[{qty}]" if qty > 1 else target_model
        if has_film: res += " +필름"
        if '택배' in s or '택배배송' in s: res += " 택배"
        return res
    
    if '케이스' not in s:
        res = s
        if has_film: res += " +필름"
        return res
    
    qty = 1
    if '☞' in s:
        parts = s.split('☞')
        base = parts[0].strip()
        after = parts[1].strip() if len(parts) > 1 else ""
        qty_match = re.match(r'^(\d+)\s*(EA|개)?', after, re.IGNORECASE)
        if qty_match: qty = int(qty_match.group(1))
    else:
        base = s
        
    if '모델선택:' in base:
        parts = base.split('모델선택:')
        target = parts[-1].strip()
        if '/' in target: target = target.split('/')[0].strip()
        if ',' in target: target = target.split(',')[0].strip()
    else:
        sub_parts = re.split(r'[/]{1,2}', base)
        target = sub_parts[-1].strip() if len(sub_parts) > 1 else base
        
    target = re.sub(r'\([^)]*\)', '', target).strip()
    target = re.sub(r'케이스', '', target)
    target = re.sub(r'풀액정', '', target)
    target = re.sub(r'2장', '', target)
    target = re.sub(r'클리어', '', target)
    target = re.sub(r'투명', '', target)
    target = re.sub(r'갤럭시', '', target)
    target = ' '.join(target.split())
    
    if target == '와이드5': target = 'A22'
    elif target == '와이드6': target = 'A13'
    
    result_str = f"{target}[{qty}]" if qty > 1 else target
    if '택배' in s or '택배배송' in s: result_str = f"{result_str} 택배"
    if has_film: result_str = f"{result_str} +필름"
        
    return result_str

# --- [ 행정안전부 주소정보누리집 API 연동 및 맞춤형 포맷 엔진 ] ---
def juso_gov_converter(keyword):
    if not keyword or pd.isna(keyword):
        return keyword
        
    kw_str = str(keyword).strip()

    # 1. 이미 도로명주소(신주소) 형태인지 판별 (구로1동 같은 동 이름 오인 방지 예외 처리 포함)
    is_already_road_address = False
    if not (re.search(r'[로길]\s*\d+\s*[동읍면]', kw_str)):
        is_already_road_address = bool(re.search(r'[로길]\s*\d+', kw_str))

    if is_already_road_address:
        parens = re.findall(r'\([^)]+\)', kw_str)
        paren_str = parens[-1] if parens else ""
        for p in parens:
            kw_str = kw_str.replace(p, "").strip()

        tail_detail = ""
        m_detail = re.search(r'(\d+층\s*\d+호|\d+동\s*\d+호|\d+동\s*\d+-\d+호|\d+동\s*\d+층|\d+-\d+호|\b\d+-\d+\b호|\d+호|\d+층)', kw_str)
        if m_detail:
            tail_detail = m_detail.group(1).strip()
            kw_str = kw_str[:m_detail.start()].strip()

        kw_str = re.sub(r'\s+', ' ', kw_str).strip()

        parts = [kw_str]
        if paren_str:
            parts.append(paren_str)
        if tail_detail:
            parts.append(tail_detail)

        final_str = ' '.join(parts)
        final_str = remove_duplicate_parentheses(final_str)
        final_str = format_dong_ho(final_str)
        return ' '.join(final_str.split())

    # --- 구주소(지번 주소)인 경우 ---
    parens = re.findall(r'\([^)]+\)', kw_str)
    paren_str = parens[-1] if parens else ""
    for p in parens:
        kw_str = kw_str.replace(p, "").strip()

    tail_detail = ""
    m_detail = re.search(r'(\d+층\s*\d+호|\d+동\s*\d+호|\d+동\s*\d+-\d+호|\d+동\s*\d+층|\d+-\d+호|\b\d+-\d+\b호|\d+호|\d+층)', kw_str)
    if m_detail:
        tail_detail = m_detail.group(1).strip()
        kw_str = kw_str[:m_detail.start()].strip()

    # 건물명 추출 및 기억 (구주소에서 아파트/빌라 명칭 분리)
    building_name = ""
    m_bldg = re.search(r'\d+(?:-\d+)?\s+([가-힣0-9]+(?:아파트|오피스텔|타워|빌딩|스퀘어|프라자|파크|시티|팰리스|마크|하우스|그랑비스타|헤리움|칸타빌|주공|빌|리슈빌|\d+차))\b', kw_str)
    if not m_bldg:
        m_bldg = re.search(r'([가-힣0-9]+(?:아파트|오피스텔|타워|빌딩|스퀘어|프라자|파크|시티|팰리스|마크|하우스|그랑비스타|헤리움|칸타빌|주공|빌|리슈빌|\d+차))\s*$', kw_str)

    if m_bldg:
        potential_bldg = m_bldg.group(1).strip()
        if len(potential_bldg) >= 2 and not re.search(r'(특별시|광역시|도|시|군|구|동|읍|면|도곡리|삼덕동\d*)$', potential_bldg):
            building_name = potential_bldg
            kw_str = kw_str.replace(building_name, "").strip()

    kw_str = re.sub(r'\s+', ' ', kw_str).strip()

    # API 요청 시에는 건물명을 제외한 깔끔한 지번 주소만 전송
    api_keyword = kw_str

    url = "https://www.juso.go.kr/addrlink/addrLinkApi.do"
    params = {
        "confmKey": JUSO_CONFIRM_KEY,
        "currentPage": 1,
        "countPerPage": 1,
        "keyword": api_keyword,
        "resultType": "json"
    }
    
    base_road_addr = kw_str
    try:
        response = requests.get(url, params=params, timeout=5)
        if response.status_code == 200:
            res = response.json()
            common = res.get('results', {}).get('common', {})
            if common.get('errorCode') == '0':
                juso_list = res.get('results', {}).get('juso', [])
                if juso_list:
                    road_addr = juso_list[0].get('roadAddr', '')
                    if road_addr:
                        base_road_addr = road_addr
    except Exception:
        pass

    # 최종 조립 순서: [도로명주소] + [건물명] + [(괄호)] + [동·호수/층]
    parts = [base_road_addr]
    if building_name and building_name not in base_road_addr:
        parts.append(building_name)
    if paren_str and paren_str not in base_road_addr:
        parts.append(paren_str)
    if tail_detail:
        parts.append(tail_detail)

    final_str = ' '.join(parts)
    final_str = remove_duplicate_parentheses(final_str)
    final_str = format_dong_ho(final_str)
    return ' '.join(final_str.split())

# --- [ Streamlit 웹 UI 실행 부문 ] ---
st.set_page_config(page_title="주소 및 상품명 변환기", page_icon="🏛️", layout="centered")

st.title("🏛️ 주소 및 상품명 엑셀 자동 정제기")
st.write("엑셀 파일을 업로드하시면 도로명주소 변환, 상세주소 우선 정렬 및 상품명 정리가 함께 수행됩니다.")

uploaded_file = st.file_uploader("변환할 엑셀 파일(.xlsx, .xls)을 업로드하세요", type=["xlsx", "xls"])

if uploaded_file is not None:
    try:
        df = pd.read_excel(uploaded_file)
    except Exception:
        df = pd.read_excel(uploaded_file, engine='xlrd')

    st.write("### 📄 업로드 데이터 미리보기 (상위 5건)")
    st.dataframe(df.head())

    if 'excel_data' not in st.session_state:
        st.session_state.excel_data = None

    if st.button("🚀 변환 및 정제 시작"):
        with st.spinner("데이터를 처리하는 중입니다. 잠시만 기다려주세요..."):
            if '우편번호' in df.columns:
                df['우편번호'] = df['우편번호'].apply(fix_zipcode)

            if '선택정보' in df.columns:
                df['선택정보'] = df['선택정보'].apply(clean_product_name)
            elif len(df.columns) >= 5:
                col_name_e = df.columns[4]
                df[col_name_e] = df[col_name_e].apply(clean_product_name)

            if '배송지' in df.columns:
                progress_bar = st.progress(0)
                total = len(df)
                
                converted_addrs = []
                for i, row in df.iterrows():
                    res = juso_gov_converter(row['배송지'])
                    converted_addrs.append(res)
                    progress_bar.progress((i + 1) / total)
                    
                df['배송지'] = converted_addrs
            else:
                st.warning("[주의] 엑셀 파일에서 '배송지' 컬럼을 찾을 수 없습니다.")

            for col in df.columns:
                df[col] = df[col].apply(remove_illegal_chars)

            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df.to_excel(writer, index=False)
                worksheet = writer.sheets['Sheet1']
                
                for col in worksheet.columns:
                    max_len = 0
                    col_letter = openpyxl.utils.get_column_letter(col[0].column)
                    for cell in col:
                        val = str(cell.value or '')
                        length = sum(2 if ord(char) > 127 else 1 for char in val)
                        if length > max_len:
                            max_len = length
                    worksheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 80)

                font_size_8 = Font(size=8)
                for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, min_col=1, max_col=len(df.columns)):
                    worksheet.row_dimensions[row[0].row].height = 18
                    for cell in row:
                        cell.font = font_size_8
                        
                if '우편번호' in df.columns:
                    zip_col_idx = df.columns.get_loc('우편번호') + 1
                    for row in range(2, worksheet.max_row + 1):
                        worksheet.cell(row=row, column=zip_col_idx).number_format = '@'

            st.session_state.excel_data = output.getvalue()
        st.success("🎉 변환 및 정리가 완료되었습니다!")

    if st.session_state.excel_data is not None:
        st.download_button(
            label="📥 변환된 엑셀 파일 다운로드",
            data=st.session_state.excel_data,
            file_name="주소_상품명_변환완료.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
